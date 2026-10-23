from __future__ import annotations

import logging
from pathlib import Path

logger = logging.getLogger(__name__)


def extract_text_from_file(path: Path) -> str:
    suffix = path.suffix.lower()
    if suffix in {".txt", ".md", ".csv"}:
        return path.read_text(encoding="utf-8", errors="ignore")
    if suffix == ".pdf":
        return _from_pdf(path)
    if suffix in {".docx"}:
        return _from_docx(path)
    if suffix in {".xlsx", ".xls"}:
        return _from_excel(path)
    if suffix == ".doc":
        raise ValueError("فرمت .doc پشتیبانی نمی‌شود؛ به .docx تبدیل کنید")
    raise ValueError(f"فرمت پشتیبانی نمی‌شود: {suffix}")


def _from_pdf(path: Path) -> str:
    try:
        from pypdf import PdfReader
    except ImportError as exc:
        raise RuntimeError("برای PDF نصب کن: pip install pypdf") from exc
    reader = PdfReader(str(path))
    parts = []
    for page in reader.pages:
        parts.append(page.extract_text() or "")
    return "\n".join(parts).strip()


def _from_docx(path: Path) -> str:
    try:
        from docx import Document
    except ImportError as exc:
        raise RuntimeError("برای Word نصب کن: pip install python-docx") from exc
    doc = Document(str(path))
    return "\n".join(p.text for p in doc.paragraphs if p.text).strip()


def _from_excel(path: Path) -> str:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise RuntimeError("برای Excel نصب کن: pip install openpyxl") from exc
    wb = load_workbook(str(path), data_only=True)
    lines: list[str] = []
    for sheet in wb.worksheets:
        lines.append(f"# Sheet: {sheet.title}")
        for row in sheet.iter_rows(values_only=True):
            cells = ["" if c is None else str(c) for c in row]
            if any(cells):
                lines.append("\t".join(cells))
    return "\n".join(lines).strip()


def analyze_transcript_text(text: str) -> dict:
    """تحلیل ساده کارنامه از متن استخراج‌شده (قبل از LLM)."""
    lines = [ln.strip() for ln in text.splitlines() if ln.strip()]
    numeric_scores: list[float] = []
    for ln in lines:
        for token in ln.replace(",", ".").split():
            try:
                val = float(token)
                if 0 <= val <= 20:
                    numeric_scores.append(val)
            except ValueError:
                continue
    avg = sum(numeric_scores) / len(numeric_scores) if numeric_scores else None
    return {
        "line_count": len(lines),
        "detected_scores": numeric_scores[:50],
        "average_score": round(avg, 2) if avg is not None else None,
        "preview": "\n".join(lines[:40]),
    }
